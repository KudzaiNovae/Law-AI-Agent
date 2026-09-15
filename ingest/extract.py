"""Extract plain text from raw documents (PDF, HTML, Word, TXT).

Scanned PDFs (common for older Gazettes/SIs) yield no text here —
those need OCR (e.g. `ocrmypdf input.pdf output.pdf`) first.
"""
import shutil
import subprocess
import tempfile
from pathlib import Path

from bs4 import BeautifulSoup


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _from_pdf(path)
    if suffix in (".html", ".htm"):
        return _from_html(path)
    if suffix in (".txt", ".md"):
        return path.read_text(encoding="utf-8", errors="replace")
    if suffix in (".docx", ".doc", ".rtf"):
        return _from_word(path)
    raise ValueError(f"Unsupported file type: {path.name}")


def _from_pdf(path: Path) -> str:
    import pdfplumber

    pages = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            pages.append(page.extract_text() or "")
    text = "\n".join(pages)
    if len(text.strip()) < 200:
        print(f"WARNING: {path.name} extracted almost no text — likely a "
              f"scanned PDF. Run OCR first (ocrmypdf).")
    return text


def _from_word(path: Path) -> str:
    """Word judgments, which is how Veritas publishes the older ones.

    Dispatch is on the first bytes rather than the extension because
    several of these are RTF or OOXML files named ".doc", and one that
    is quietly misread as the wrong format loses a whole judgment.
    """
    head = path.read_bytes()[:8]
    if head[:2] == b"PK":                     # OOXML is a zip
        return _from_docx(path)
    return _from_legacy_word(path)            # RTF or OLE2


def _from_docx(path: Path) -> str:
    import docx

    d = docx.Document(str(path))
    blocks = [p.text for p in d.paragraphs]
    # Judgments put the order, and sometimes the headnote, in a table.
    for table in d.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                blocks.append(" | ".join(cells))
    return "\n".join(b for b in blocks if b.strip())


def _from_legacy_word(path: Path) -> str:
    """Binary .doc and RTF, via LibreOffice.

    Only ingest needs this, and ingest runs on a workstation rather than
    in the container, so the dependency never reaches the image. It is
    still optional: a missing LibreOffice raises like any other
    extraction failure and run.py carries on with the next file.
    """
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        raise RuntimeError(
            f"{path.name} is a legacy Word file and LibreOffice is not "
            f"installed (apt install libreoffice-writer)")
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run(
            [soffice, "--headless",
             # our own profile: LibreOffice refuses to start a second
             # instance against a profile the desktop session is using
             f"-env:UserInstallation=file://{tmp}/profile",
             "--convert-to", "txt:Text (encoded):UTF8",
             "--outdir", tmp, str(path)],
            check=True, capture_output=True, timeout=180)
        produced = list(Path(tmp).glob("*.txt"))
        if not produced:
            raise RuntimeError(f"LibreOffice produced no text for {path.name}")
        return produced[0].read_text(encoding="utf-8", errors="replace")


def _from_html(path: Path) -> str:
    soup = BeautifulSoup(path.read_text(encoding="utf-8", errors="replace"),
                         "html.parser")
    # ZimLII Akoma Ntoso: main content lives in <div class="akn-akomaNtoso">
    main = (soup.find(class_="akn-akomaNtoso")
            or soup.find("main")
            or soup.find("article")
            or soup.body
            or soup)
    for tag in main.find_all(["script", "style", "nav", "header", "footer"]):
        tag.decompose()
    return main.get_text(separator="\n", strip=True)
